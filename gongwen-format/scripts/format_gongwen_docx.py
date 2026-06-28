#!/usr/bin/env python3
"""Format a .docx or legacy .doc file as a standard company official document.

Flow implemented:
1. Load rule library.
2. Parse Word structure.
3. Identify document semantics.
4. Apply formatting and produce a review report.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RULES = ROOT / "references" / "gongwen_rules.json"


@dataclass
class Finding:
    kind: str
    paragraph: int | None
    text: str
    action: str
    status: str = "已处理"


@dataclass
class ParagraphRole:
    index: int
    text: str
    role: str
    level: int | None = None


@dataclass
class FormatContext:
    rules: dict
    roles: list[ParagraphRole] = field(default_factory=list)
    findings: list[Finding] = field(default_factory=list)


def load_rules(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def clean_text(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def visible_paragraphs(doc: Document) -> list[tuple[int, str]]:
    items = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            items.append((i, text))
    return items


def detect_roles(doc: Document, rules: dict) -> list[ParagraphRole]:
    patterns = rules["semantic_patterns"]
    visible = visible_paragraphs(doc)
    roles: list[ParagraphRole] = []
    title_idx = visible[0][0] if visible else None

    date_candidates = {
        idx
        for idx, text in visible
        if re.match(patterns["date"], clean_text(text))
    }
    last_date_idx = max(date_candidates) if date_candidates else None

    for order, (idx, text) in enumerate(visible):
        stripped = text.strip()
        compact = clean_text(stripped)
        role = "body"
        level = None

        if idx == title_idx:
            role = "title"
        elif order <= 4 and re.search(patterns["main_receiver"], stripped):
            role = "main_receiver"
        elif re.match(patterns["attachment_note"], compact):
            role = "attachment_note"
        elif idx == last_date_idx:
            role = "date"
        elif last_date_idx is not None and idx == last_date_idx - 1:
            role = "issuer"
        elif re.match(patterns["level_1_heading"], compact):
            role, level = "heading", 1
        elif re.match(patterns["level_2_heading"], compact):
            role, level = "heading", 2
        elif re.match(patterns["level_3_heading"], compact):
            role, level = "heading", 3
        elif re.match(patterns["level_4_heading"], compact):
            role, level = "heading", 4

        roles.append(ParagraphRole(idx, stripped, role, level))

    return roles


def iter_runs(paragraph):
    if paragraph.runs:
        yield from paragraph.runs
    else:
        paragraph.add_run("")
        yield from paragraph.runs


def set_run_font(run, name: str, size_pt: int, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size_pt)
    run.bold = bold


def set_paragraph_spacing(paragraph, line_pt: int, first_line_chars: int | None = None) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(line_pt)
    if first_line_chars is None:
        fmt.first_line_indent = None
    else:
        # A 3号 Chinese character is 16pt wide; two chars = 32pt.
        fmt.first_line_indent = Pt(16 * first_line_chars)


def set_paragraph_bool_property(paragraph, tag: str, enabled: bool) -> None:
    """Set a paragraph boolean property directly to override inherited template styles."""
    ppr = paragraph._p.get_or_add_pPr()
    elem = ppr.find(qn(f"w:{tag}"))
    if elem is None:
        elem = OxmlElement(f"w:{tag}")
        ppr.append(elem)
    elem.set(qn("w:val"), "1" if enabled else "0")


def suppress_template_pagination(doc: Document, ctx: FormatContext) -> None:
    """Disable inherited template pagination that can create blank pages."""
    for paragraph in doc.paragraphs:
        set_paragraph_bool_property(paragraph, "pageBreakBefore", False)
        set_paragraph_bool_property(paragraph, "keepNext", False)
        set_paragraph_bool_property(paragraph, "keepLines", False)
    add_finding_once(
        ctx,
        Finding("格式修正", None, "分页控制", "已覆盖模板样式中的段前分页、与下段同页、段中不分页，避免生成空白页"),
    )


def paragraph_font_snapshot(paragraph) -> str:
    parts = []
    for run in paragraph.runs:
        parts.append(
            f"{run.font.name or ''}/{run.font.size.pt if run.font.size else ''}/"
            f"{bool(run.bold)}"
        )
    return ";".join(parts)


def replace_paragraph_text(paragraph, new_text: str) -> None:
    paragraph.text = new_text


def alignment_from_rule(value: str):
    mapping = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    return mapping.get(value, WD_ALIGN_PARAGRAPH.LEFT)


def add_finding_once(ctx: FormatContext, finding: Finding) -> None:
    key = (finding.kind, finding.paragraph, finding.text, finding.action)
    for item in ctx.findings:
        if (item.kind, item.paragraph, item.text, item.action) == key:
            return
    ctx.findings.append(finding)


def apply_format_error_rules(doc: Document, ctx: FormatContext) -> None:
    """Detect PPT '避开雷区' formatting examples and auto-fix safe cases."""
    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        new_text = text
        para_no = i + 1

        if para_no == 1 and re.search(r"[，,。；;：:]$", new_text):
            fixed = re.sub(r"[，,。；;：:]+$", "", new_text)
            add_finding_once(ctx, Finding("自动修正", para_no, text, "删除公文标题末尾标点"))
            new_text = fixed

        if re.match(r"^[\[〔]\d{4}[\]〕]\d{1,2}月\d{1,2}日$", clean_text(new_text)):
            fixed = re.sub(r"^[\[〔](\d{4})[\]〕](\d{1,2}月\d{1,2}日)$", r"\1年\2", clean_text(new_text))
            add_finding_once(ctx, Finding("自动修正", para_no, text, "成文日期由括号年份修正为“YYYY年M月D日”"))
            new_text = fixed

        if "》、《" in new_text or "”、“" in new_text:
            fixed = new_text.replace("》、《", "》《").replace("”、“", "”“")
            add_finding_once(ctx, Finding("自动修正", para_no, text, "删除多个书名号或引号并列时的顿号"))
            new_text = fixed

        if re.search(r"^\d+、", clean_text(new_text)):
            fixed = re.sub(r"^(\d+)、", r"\1.", new_text)
            add_finding_once(ctx, Finding("自动修正", para_no, text, "阿拉伯数字层级序号由“1、”修正为“1.”"))
            new_text = fixed

        if re.search(r"\[\d{4}\]", new_text):
            fixed = re.sub(r"\[(\d{4})\]", r"〔\1〕", new_text)
            add_finding_once(ctx, Finding("自动修正", para_no, text, "发文年号由方括号修正为六角括号“〔〕”"))
            new_text = fixed

        if re.search(r"\d{4}-\d{4}年", new_text):
            fixed = re.sub(r"(\d{4})-(\d{4}年)", r"\1—\2", new_text)
            add_finding_once(ctx, Finding("自动修正", para_no, text, "年份起止连接号由“-”修正为一字线“—”"))
            new_text = fixed

        if re.match(r"^附件\d+[:：]", clean_text(new_text)):
            fixed = re.sub(r"^附件(\d+)[:：]\s*", r"附件：\1.", new_text)
            add_finding_once(ctx, Finding("自动修正", para_no, text, "附件序号由“附件N：”修正为“附件：N.”"))
            new_text = fixed

        if re.match(r"^附件[:：]", clean_text(new_text)) and re.search(r"[。；;,.，、]$", new_text):
            fixed = re.sub(r"[。；;,.，、]+$", "", new_text)
            add_finding_once(ctx, Finding("自动修正", para_no, text, "删除附件名称末尾标点"))
            new_text = fixed

        if re.match(r"^（[一二三四五六七八九十]+）", clean_text(new_text)) and re.search(r"[。.]$", new_text):
            fixed = re.sub(r"[。.]$", "", new_text)
            add_finding_once(ctx, Finding("自动修正", para_no, text, "二级标题独立成段时删除末尾句号"))
            new_text = fixed

        if re.match(r"^注[:：]", clean_text(new_text)) and re.search(r"[。.]$", new_text):
            fixed = re.sub(r"[。.]$", "", new_text)
            add_finding_once(ctx, Finding("自动修正", para_no, text, "图表注释类说明文字删除末尾句号"))
            new_text = fixed

        if re.search(r"（[^（）]*（[^（）]*）[^（）]*）", new_text):
            add_finding_once(ctx, Finding("需复核", para_no, text, "存在同一形式括号套用，建议改用不同括号形式配合", "需人工确认"))

        if new_text != text:
            replace_paragraph_text(paragraph, new_text)


def apply_paragraph_style(doc: Document, role: ParagraphRole, ctx: FormatContext) -> None:
    rules = ctx.rules
    p = doc.paragraphs[role.index]
    before = paragraph_font_snapshot(p)
    text = role.text
    line_pt = rules["paragraph"]["line_spacing_pt"]
    body_font = rules["fonts"]["body"]

    if role.role == "title":
        cfg = rules["fonts"]["title"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, line_pt, None)
        for run in iter_runs(p):
            set_run_font(run, cfg["name"], cfg["size_pt"], cfg.get("bold", False))
        ctx.findings.append(Finding("格式修正", role.index + 1, text, "按标题设置为2号方正小标宋体、居中、不加粗"))
        return

    if role.role == "main_receiver":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_paragraph_spacing(p, line_pt, None)
    elif role.role in {"issuer", "date"}:
        signature = rules.get("signature", {})
        key = "issuer_alignment" if role.role == "issuer" else "date_alignment"
        p.alignment = alignment_from_rule(signature.get(key, "right"))
        set_paragraph_spacing(p, line_pt, None)
    elif role.role == "attachment_note":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_paragraph_spacing(p, line_pt, rules["paragraph"]["body_first_line_chars"])
        if re.search(r"[。；;,.，、]$", text):
            ctx.findings.append(Finding("需复核", role.index + 1, text, "附件名称后不应加标点，已提示人工确认", "需人工确认"))
    elif role.role == "heading":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_paragraph_spacing(p, line_pt, rules["paragraph"]["body_first_line_chars"])
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_paragraph_spacing(p, line_pt, rules["paragraph"]["body_first_line_chars"])

    if role.role == "heading" and role.level == 1:
        cfg = rules["fonts"]["level_1_heading"]
        action = "按一级标题设置为3号黑体、不加粗"
    elif role.role == "heading" and role.level == 2:
        cfg = rules["fonts"]["level_2_heading"]
        action = "按二级标题设置为3号楷体_GB2312、不加粗"
    elif role.role == "heading" and role.level in {3, 4}:
        cfg = rules["fonts"]["level_3_heading"]
        action = "按三级/四级标题设置为3号仿宋_GB2312"
    elif role.role == "issuer":
        cfg = body_font
        action = "按落款设置为3号仿宋_GB2312、右对齐"
    elif role.role == "date":
        cfg = body_font
        action = "按成文日期设置为3号仿宋_GB2312、右对齐"
    else:
        cfg = body_font
        action = "按正文设置为3号仿宋_GB2312、固定值28磅行距"

    for run in iter_runs(p):
        set_run_font(run, cfg["name"], cfg["size_pt"], cfg.get("bold", False))

    after = paragraph_font_snapshot(p)
    if before != after or role.role != "body":
        ctx.findings.append(Finding("格式修正", role.index + 1, text[:80], action))


def apply_page_setup(doc: Document, ctx: FormatContext) -> None:
    page = ctx.rules["page"]
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(page["width_mm"])
        section.page_height = Mm(page["height_mm"])
        section.top_margin = Mm(page["top_mm"])
        section.left_margin = Mm(page["left_mm"])
        section.right_margin = Mm(page["right_mm"])
        section.bottom_margin = Mm(page["bottom_mm"])
    ctx.findings.append(Finding("格式修正", None, "页面设置", "设置为A4，天头37mm，订口28mm，右边距26mm，下边距35mm"))


def add_page_number(section, rules: dict) -> None:
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ""
    font = rules["fonts"]["footer_page_number"]
    for text in ("-", " "):
        r = p.add_run(text)
        set_run_font(r, font["name"], font["size_pt"], False)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = p.add_run()
    set_run_font(r, font["name"], font["size_pt"], False)
    r._r.append(fld_begin)
    r._r.append(instr)
    r._r.append(fld_sep)
    r._r.append(text)
    r._r.append(fld_end)
    for text in (" ", "-"):
        r = p.add_run(text)
        set_run_font(r, font["name"], font["size_pt"], False)


def apply_footers(doc: Document, ctx: FormatContext) -> None:
    for section in doc.sections:
        add_page_number(section, ctx.rules)
    ctx.findings.append(Finding("格式修正", None, "页码", "设置页脚居中页码样式：- PAGE -"))


def add_semantic_findings(ctx: FormatContext) -> None:
    roles = ctx.roles
    title = next((r for r in roles if r.role == "title"), None)
    if title:
        if re.search(r"[，,。；;：:]", title.text):
            ctx.findings.append(Finding("需复核", title.index + 1, title.text, "公文标题除法规、规章名称加书名号外，一般不用标点", "需人工确认"))
        if title.text.count("关于") > 1 or title.text.count("通知") > 1:
            ctx.findings.append(Finding("需复核", title.index + 1, title.text, "标题可能存在介词或文种重叠", "需人工确认"))
    if not any(r.role == "main_receiver" for r in roles):
        ctx.findings.append(Finding("需复核", None, "主送机关", "未识别到标题下方主送机关，请人工确认", "需人工确认"))
    if not any(r.role == "date" for r in roles):
        ctx.findings.append(Finding("需复核", None, "成文日期", "未识别到形如“2026年6月6日”的成文日期", "需人工确认"))
    for r in roles:
        if re.search(r"^\d+、", clean_text(r.text)):
            add_finding_once(ctx, Finding("需复核", r.index + 1, r.text, "阿拉伯数字序号建议使用“1.”，不使用“1、”", "需人工确认"))
        if re.search(r"\(\d{4}-\d{4}年\)", r.text):
            add_finding_once(ctx, Finding("需复核", r.index + 1, r.text, "年份起止建议使用一字线，如“2013—2015年”", "需人工确认"))


def format_document(input_path: Path, output_path: Path, report_path: Path, rules_path: Path) -> FormatContext:
    rules = load_rules(rules_path)
    doc = Document(str(input_path))
    ctx = FormatContext(rules=rules)
    ctx.roles = detect_roles(doc, rules)

    apply_page_setup(doc, ctx)
    suppress_template_pagination(doc, ctx)
    apply_format_error_rules(doc, ctx)
    ctx.roles = detect_roles(doc, rules)
    add_semantic_findings(ctx)
    for role in ctx.roles:
        apply_paragraph_style(doc, role, ctx)
    apply_footers(doc, ctx)

    doc.save(str(output_path))
    write_report(report_path, input_path, output_path, ctx)
    return ctx


def write_report(path: Path, input_path: Path, output_path: Path, ctx: FormatContext) -> None:
    doc = Document()
    doc.add_heading("公文格式智能核对报告", 0)
    doc.add_paragraph(f"原始文件：{input_path.name}")
    doc.add_paragraph(f"规范化文件：{output_path.name}")
    doc.add_paragraph(f"规则库版本：{ctx.rules.get('version', '')}")
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    doc.add_heading("一、识别结果", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "段落"
    hdr[1].text = "角色"
    hdr[2].text = "层级"
    hdr[3].text = "文本摘录"
    for r in ctx.roles:
        row = table.add_row().cells
        row[0].text = str(r.index + 1)
        row[1].text = r.role
        row[2].text = "" if r.level is None else str(r.level)
        row[3].text = r.text[:90]

    doc.add_heading("二、核对与处理清单", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "类型"
    hdr[1].text = "段落"
    hdr[2].text = "对象"
    hdr[3].text = "处理/提示"
    hdr[4].text = "状态"
    for f in ctx.findings:
        row = table.add_row().cells
        row[0].text = f.kind
        row[1].text = "" if f.paragraph is None else str(f.paragraph)
        row[2].text = f.text
        row[3].text = f.action
        row[4].text = f.status

    doc.add_heading("三、仍需人工复核", level=1)
    for item in ctx.rules.get("manual_review", []):
        doc.add_paragraph(item, style=None)

    for p in doc.paragraphs:
        for run in p.runs:
            set_run_font(run, "仿宋_GB2312", 12 if p.style.name != "Title" else 18, bool(run.bold))

    doc.save(str(path))


def ensure_supported_document(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    if path.suffix.lower() not in {".docx", ".doc"}:
        raise ValueError("仅支持 .docx 或 .doc 文件")


def convert_doc_to_docx(input_path: Path, output_dir: Path) -> Path:
    """Convert legacy .doc to .docx using an available local converter."""
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{input_path.stem}.docx"

    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office:
        subprocess.run(
            [office, "--headless", "--convert-to", "docx", "--outdir", str(output_dir), str(input_path)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        converted = output_dir / f"{input_path.stem}.docx"
        if converted.exists():
            return converted

    textutil = shutil.which("textutil")
    if textutil:
        subprocess.run(
            [textutil, "-convert", "docx", "-output", str(output_path), str(input_path)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        if output_path.exists():
            return output_path

    raise RuntimeError("未找到 .doc 转 .docx 工具。请安装 LibreOffice/soffice，或在 macOS 使用 textutil。")


def prepare_input_document(input_path: Path, output_dir: Path) -> tuple[Path, bool]:
    ensure_supported_document(input_path)
    if input_path.suffix.lower() == ".docx":
        return input_path, False
    converted_dir = output_dir / "_converted"
    return convert_doc_to_docx(input_path, converted_dir), True


def default_paths(input_path: Path, output_dir: Path) -> tuple[Path, Path]:
    stem = input_path.stem
    return (
        output_dir / f"{stem}-标准公文格式.docx",
        output_dir / f"{stem}-格式核对报告.docx",
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="公文格式智能核对与规范化助手")
    parser.add_argument("input", type=Path, help="待规范化的 .docx 或 .doc 文件")
    parser.add_argument("-o", "--output-dir", type=Path, default=Path("output"), help="输出目录")
    parser.add_argument("--rules", type=Path, default=DEFAULT_RULES, help="规则库 JSON")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    output_path, report_path = default_paths(args.input, args.output_dir)
    prepared_input, converted = prepare_input_document(args.input, args.output_dir)
    ctx = format_document(prepared_input, output_path, report_path, args.rules)
    if converted:
        ctx.findings.insert(0, Finding("格式转换", None, args.input.name, "已先将 .doc 转换为 .docx 后再进行公文格式处理"))
        write_report(report_path, args.input, output_path, ctx)
    print(f"规范化文档：{output_path}")
    print(f"核对报告：{report_path}")
    print(f"识别段落：{len(ctx.roles)}")
    print(f"处理/提示项：{len(ctx.findings)}")


if __name__ == "__main__":
    main()
